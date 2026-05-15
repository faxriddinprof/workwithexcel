from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


OUTPUT_PATH = Path("media/report_templates/2026/05/hard_formula_test_template_2026_05.xlsx")
ROW_COUNT = 180


def build_rules_sheet(workbook):
    worksheet = workbook.create_sheet("Rules")
    worksheet.append(["Tier", "RevenueMultiplier", "QualityMultiplier", "QualityBand", "NetAdjust"])
    tier_rows = [
        ["Tier-1", 1.55, None, None, 1.18],
        ["Tier-2", 1.35, None, None, 1.12],
        ["Tier-3", 1.22, None, None, 1.07],
        ["Tier-4", 1.10, None, None, 1.03],
        ["Tier-5", 0.95, None, None, 0.98],
    ]
    for row in tier_rows:
        worksheet.append(row)

    worksheet["C1"] = "QualityMultiplier"
    worksheet["D1"] = "Band"
    worksheet["C2"] = 1.25
    worksheet["D2"] = "A"
    worksheet["C3"] = 1.10
    worksheet["D3"] = "B"
    worksheet["C4"] = 0.92
    worksheet["D4"] = "C"
    worksheet["C5"] = 0.75
    worksheet["D5"] = "D"

    worksheet["G1"] = "Region"
    worksheet["H1"] = "AuditGroup"
    region_rules = [
        ["Tashkent", "Priority"],
        ["Samarkand", "Priority"],
        ["Bukhara", "Priority"],
        ["Andijan", "Standard"],
        ["Namangan", "Standard"],
        ["Fergana", "Standard"],
        ["Kashkadarya", "Priority"],
        ["Surkhandarya", "Standard"],
        ["Khorezm", "Standard"],
        ["Jizzakh", "Standard"],
        ["Navoi", "Priority"],
        ["Syrdarya", "Standard"],
    ]
    for index, (region, group) in enumerate(region_rules, start=2):
        worksheet[f"G{index}"] = region
        worksheet[f"H{index}"] = group

    worksheet["I1"] = "MonthStart"
    worksheet["J1"] = "SeasonFactor"
    seasonal_rules = [
        [1, 0.95],
        [4, 1.02],
        [7, 1.09],
        [10, 1.04],
    ]
    for index, (month_start, factor) in enumerate(seasonal_rules, start=2):
        worksheet[f"I{index}"] = month_start
        worksheet[f"J{index}"] = factor

    for column in range(1, 11):
        worksheet.column_dimensions[get_column_letter(column)].width = 18


def build_data_sheet(workbook):
    worksheet = workbook.active
    worksheet.title = "Data"
    headers = [
        "Region",
        "District",
        "Specialist",
        "PlannedUnits",
        "ActualUnits",
        "UnitPrice",
        "DaysWorked",
        "OvertimeHours",
        "QualityScore",
        "CompletionPct",
        "ReportDate",
        "TransportCost",
        "MealCost",
        "BonusRate",
        "LookupTier",
        "Variance",
        "VariancePct",
        "CostBase",
        "OvertimeCost",
        "TotalCost",
        "Efficiency",
        "QualityBand",
        "PerfScore",
        "PerfBonus",
        "RiskFlag",
        "RevenueEstimate",
        "NetResult",
        "NetMargin",
        "TextKey",
        "LookupMultiplier",
        "AdjustedNet",
        "RankScore",
        "HealthCheck",
        "AuditNote",
        "WeightedIndex",
        "SeasonAdjustedNet",
    ]
    worksheet.append(headers)

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    formula_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(color="FFFFFF", bold=True)

    for column_index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column_index)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(14, len(header) + 2)
        if column_index >= 16:
            worksheet.cell(row=2, column=column_index).fill = formula_fill

    regions = [
        "Tashkent",
        "Samarkand",
        "Bukhara",
        "Andijan",
        "Namangan",
        "Fergana",
        "Kashkadarya",
        "Surkhandarya",
        "Khorezm",
        "Jizzakh",
        "Navoi",
        "Syrdarya",
    ]
    districts = [
        "Central",
        "North",
        "South",
        "East",
        "West",
        "Industrial",
        "Rural",
        "Urban",
    ]
    specialists = [
        "Aziza Karimova",
        "Bekzod Aliyev",
        "Dilshod Rustamov",
        "Malika Ergasheva",
        "Shahzod Islomov",
        "Nilufar Tursunova",
        "Sardor Jabborov",
        "Kamola Rahimova",
    ]
    tiers = ["Tier-1", "Tier-2", "Tier-3", "Tier-4", "Tier-5"]

    for row in range(2, ROW_COUNT + 2):
        seed = row - 1
        region = regions[seed % len(regions)]
        district = districts[seed % len(districts)]
        specialist = specialists[seed % len(specialists)]
        planned_units = 80 + (seed * 7) % 170
        actual_units = planned_units + ((seed % 9) - 4) * 6 + (seed % 5)
        unit_price = 115 + (seed % 14) * 9
        days_worked = 18 + (seed % 9)
        overtime_hours = (seed * 3) % 22
        quality_score = 68 + (seed * 5) % 33
        completion_pct = round(min(1.35, max(0.52, actual_units / planned_units)), 4)
        report_month = ((seed - 1) % 12) + 1
        transport_cost = 40 + (seed % 12) * 8
        meal_cost = 20 + (seed % 10) * 5
        bonus_rate = round(0.03 + (seed % 6) * 0.0125, 4)
        lookup_tier = tiers[seed % len(tiers)]

        worksheet.cell(row=row, column=1, value=region)
        worksheet.cell(row=row, column=2, value=f"{district}-{(seed % 4) + 1}")
        worksheet.cell(row=row, column=3, value=specialist)
        worksheet.cell(row=row, column=4, value=planned_units)
        worksheet.cell(row=row, column=5, value=actual_units)
        worksheet.cell(row=row, column=6, value=unit_price)
        worksheet.cell(row=row, column=7, value=days_worked)
        worksheet.cell(row=row, column=8, value=overtime_hours)
        worksheet.cell(row=row, column=9, value=quality_score)
        worksheet.cell(row=row, column=10, value=completion_pct)
        worksheet.cell(row=row, column=11, value=f"2026-{report_month:02d}-01")
        worksheet.cell(row=row, column=12, value=transport_cost)
        worksheet.cell(row=row, column=13, value=meal_cost)
        worksheet.cell(row=row, column=14, value=bonus_rate)
        worksheet.cell(row=row, column=15, value=lookup_tier)

        worksheet.cell(row=row, column=16, value=f"=E{row}-D{row}")
        worksheet.cell(row=row, column=17, value=f"=IF(D{row}=0,0,E{row}/D{row}-1)")
        worksheet.cell(row=row, column=18, value=f"=E{row}*F{row}")
        worksheet.cell(row=row, column=19, value=f"=H{row}*(F{row}*1.5/8)")
        worksheet.cell(row=row, column=20, value=f"=R{row}+S{row}+L{row}+M{row}")
        worksheet.cell(row=row, column=21, value=f"=IF(G{row}=0,0,E{row}/G{row})")
        worksheet.cell(row=row, column=22, value=f"=IF(I{row}>=95,\"A\",IF(I{row}>=85,\"B\",IF(I{row}>=75,\"C\",\"D\")))")
        worksheet.cell(row=row, column=23, value=f"=ROUND((J{row}*35)+(I{row}*0.25)+(U{row}*0.2)+(IF(Q{row}>0,100,MAX(0,100+Q{row}*100))*0.2),2)")
        worksheet.cell(row=row, column=24, value=f"=ROUND(T{row}*N{row}*IF(V{row}=\"A\",1.25,IF(V{row}=\"B\",1.1,IF(V{row}=\"C\",0.9,0.5))),2)")
        worksheet.cell(row=row, column=25, value=f"=IF(OR(I{row}<80,J{row}<0.75,E{row}<D{row}*0.85),\"HIGH\",IF(OR(I{row}<90,J{row}<0.9),\"MEDIUM\",\"LOW\"))")
        worksheet.cell(row=row, column=26, value=f"=ROUND(E{row}*INDEX(Rules!$B$2:$B$6,MATCH(O{row},Rules!$A$2:$A$6,0))*INDEX(Rules!$C$2:$C$5,MATCH(V{row},Rules!$D$2:$D$5,0)),2)")
        worksheet.cell(row=row, column=27, value=f"=Z{row}-T{row}-X{row}")
        worksheet.cell(row=row, column=28, value=f"=IF(Z{row}=0,0,AA{row}/Z{row})")
        worksheet.cell(row=row, column=29, value=f"=UPPER(LEFT(A{row},3)&\"-\"&LEFT(B{row},3)&\"-\"&TEXT(ROW()-1,\"000\"))")
        worksheet.cell(row=row, column=30, value=f"=INDEX(Rules!$E$2:$E$6,MATCH(O{row},Rules!$A$2:$A$6,0))")
        worksheet.cell(row=row, column=31, value=f"=ROUND(AA{row}*AD{row},2)")
        worksheet.cell(row=row, column=32, value=f"=RANK(W{row},$W$2:$W${ROW_COUNT + 1},0)")
        worksheet.cell(row=row, column=33, value=f"=IF(AND(AB{row}>0.18,Y{row}=\"LOW\"),\"PASS\",IF(AB{row}>0.08,\"WATCH\",\"FAIL\"))")
        worksheet.cell(row=row, column=34, value=f"=IFERROR(IF(VLOOKUP(A{row},Rules!$G$2:$H$13,2,FALSE)=\"Priority\",\"Priority Region\",\"Standard Region\"),\"Standard Region\")")
        worksheet.cell(row=row, column=35, value=f"=ROUND(SUMPRODUCT(E{row}:J{row},{{0.08,0.08,0.15,0.1,0.25,0.34}}),2)")
        worksheet.cell(row=row, column=36, value=f"=ROUND(AE{row}*INDEX(Rules!$J$2:$J$5,MATCH(MONTH(K{row}),Rules!$I$2:$I$5,1)),2)")

    for row in range(2, ROW_COUNT + 2):
        for column in range(6, 37):
            cell = worksheet.cell(row=row, column=column)
            if column in {10, 14, 17, 28}:
                cell.number_format = "0.00%"
            elif column in {6, 12, 13, 18, 19, 20, 24, 26, 27, 31, 36}:
                cell.number_format = '#,##0.00'
            elif column == 11:
                cell.number_format = "yyyy-mm-dd"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:AJ{ROW_COUNT + 1}"


def build_summary_sheet(workbook):
    worksheet = workbook.create_sheet("Summary")
    rows = [
        ["Metric", "Formula"],
        ["Total Rows", "=COUNTA(Data!A2:A181)"],
        ["Total Actual Units", "=SUM(Data!E2:E181)"],
        ["Total Cost", "=SUM(Data!T2:T181)"],
        ["Total Revenue", "=SUM(Data!Z2:Z181)"],
        ["Average Margin", "=AVERAGE(Data!AB2:AB181)"],
        ["High Risk Count", "=COUNTIF(Data!Y2:Y181,\"HIGH\")"],
        ["Pass Health Count", "=COUNTIF(Data!AG2:AG181,\"PASS\")"],
        ["Priority Region Count", "=COUNTIF(Data!AH2:AH181,\"Priority Region\")"],
        ["Top Performance Score", "=MAX(Data!W2:W181)"],
        ["Lowest Net Result", "=MIN(Data!AA2:AA181)"],
        ["Season Adjusted Total", "=SUM(Data!AJ2:AJ181)"],
    ]
    for row in rows:
        worksheet.append(row)

    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 24
    worksheet.freeze_panes = "A2"


def main():
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    build_data_sheet(workbook)
    build_rules_sheet(workbook)
    build_summary_sheet(workbook)

    workbook.save(OUTPUT_PATH)
    print(f"Created: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()