from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter


def extract_formula_rules(worksheet, columns, data_start_row, scan_limit=50):
    formula_rules = []
    max_row = min(worksheet.max_row, data_start_row + scan_limit)

    for column in columns:
        for row_index in range(data_start_row, max_row + 1):
            cell = worksheet.cell(row=row_index, column=column['index'])
            if isinstance(cell.value, str) and cell.value.startswith('='):
                formula_rules.append(
                    {
                        'key': column['key'],
                        'column_index': column['index'],
                        'source_cell': cell.coordinate,
                        'source_row': row_index,
                        'formula': cell.value,
                    }
                )
                break

    formula_keys = {rule['key'] for rule in formula_rules}
    for column in columns:
        column['is_formula'] = column['key'] in formula_keys
        column['read_only'] = column['is_formula']

    return formula_rules


def build_formula_map(structure):
    return {rule['key']: rule for rule in structure.get('formulas', [])}


def translated_formula(rule, target_row, target_column=None):
    target_column = target_column or rule['column_index']
    target_cell = f"{get_column_letter(target_column)}{target_row}"
    return Translator(rule['formula'], origin=rule['source_cell']).translate_formula(target_cell)


def formula_values_for_row(structure, target_row):
    formula_map = build_formula_map(structure)
    return {
        key: translated_formula(rule, target_row, target_column=rule['column_index'])
        for key, rule in formula_map.items()
    }