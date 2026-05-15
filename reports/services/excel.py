from io import BytesIO
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import UploadedFile
from django.utils.text import slugify

from .formulas import build_formula_map, extract_formula_rules, formula_values_for_grid_row, formula_values_for_row, translated_formula


def parse_template_structure(excel_file):
    workbook = _load_workbook(excel_file)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    header_index, headers = _detect_headers(rows)
    data_start_row = header_index + 2
    columns = []
    used_labels = set()

    for index, label in enumerate(headers, start=1):
        normalized = str(label).strip() if label is not None else f'Column {index}'
        key = slugify(normalized) or f'column_{index}'
        while key in used_labels:
            key = f'{key}_{index}'
        used_labels.add(key)
        columns.append(
            {
                'key': key,
                'label': normalized,
                'index': index,
                'is_formula': False,
                'read_only': False,
            }
        )

    formula_rules = extract_formula_rules(worksheet, columns, data_start_row)

    preview_rows = []
    for raw_row in rows[header_index + 1: header_index + 4]:
        if not any(value not in (None, '') for value in raw_row[: len(columns)]):
            continue
        preview_rows.append(
            {
                column['key']: _serialize_cell(raw_row[position])
                for position, column in enumerate(columns)
            }
        )

    return {
        'sheet_name': worksheet.title,
        'header_row': header_index + 1,
        'data_start_row': data_start_row,
        'columns': columns,
        'formulas': formula_rules,
        'preview_rows': preview_rows,
    }


def rows_for_grid(submitted_data, structure):
    columns = structure.get('columns', [])
    data = []
    for row_offset, row in enumerate(submitted_data or []):
        grid_row = {column['key']: row.get(column['key'], '') for column in columns if not column.get('is_formula')}
        grid_row.update(formula_values_for_grid_row(structure, row_offset + 1))
        data.append(grid_row)
    return data or [blank_row(structure, row_offset=index) for index in range(5)]


def blank_row(structure, row_offset=0):
    row = {column['key']: '' for column in structure.get('columns', []) if not column.get('is_formula')}
    row.update(formula_values_for_grid_row(structure, row_offset + 1))
    return row


def normalize_submission_rows(raw_rows, structure):
    columns = structure.get('columns', [])
    editable_columns = [column for column in columns if not column.get('is_formula')]
    if not isinstance(raw_rows, list):
        raise ValidationError('Invalid spreadsheet payload.')

    normalized_rows = []
    for raw_row in raw_rows:
        row = {}
        if isinstance(raw_row, dict):
            source = raw_row
        elif isinstance(raw_row, list):
            source = {column['key']: raw_row[index] if index < len(raw_row) else '' for index, column in enumerate(columns)}
        else:
            raise ValidationError('Rows must be dictionaries or arrays.')

        has_value = False
        for column in editable_columns:
            value = source.get(column['key'], '')
            if value is None:
                value = ''
            if isinstance(value, str):
                value = value.strip()
            if value != '':
                has_value = True
            row[column['key']] = value
        if has_value:
            normalized_rows.append(row)
    return normalized_rows


def export_submission_workbook(report_template, submission_rows):
    workbook = _load_workbook(report_template.excel_file)
    structure = report_template.structure
    worksheet = workbook[structure['sheet_name']]
    columns = structure.get('columns', [])
    formula_map = build_formula_map(structure)
    header_row = structure.get('header_row', 1)
    data_start_row = structure.get('data_start_row', header_row + 1)

    for index, column in enumerate(columns, start=1):
        worksheet.cell(row=header_row, column=index, value=column['label'])

    for row_index in range(data_start_row, max(worksheet.max_row, data_start_row + len(submission_rows)) + 1):
        for column_index in range(1, len(columns) + 1):
            worksheet.cell(row=row_index, column=column_index, value=None)

    for row_offset, row in enumerate(submission_rows, start=data_start_row):
        for column_index, column in enumerate(columns, start=1):
            rule = formula_map.get(column['key'])
            if rule:
                worksheet.cell(
                    row=row_offset,
                    column=column_index,
                    value=translated_formula(rule, row_offset, target_column=column_index),
                )
            else:
                worksheet.cell(row=row_offset, column=column_index, value=row.get(column['key'], ''))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _load_workbook(excel_file):
    from openpyxl import load_workbook

    try:
        if isinstance(excel_file, (str, Path)):
            with open(excel_file, 'rb') as workbook_file:
                return load_workbook(workbook_file)

        if isinstance(excel_file, UploadedFile):
            if hasattr(excel_file, 'seek'):
                excel_file.seek(0)
            workbook = load_workbook(excel_file)
            if hasattr(excel_file, 'seek'):
                excel_file.seek(0)
            return workbook

        if hasattr(excel_file, 'open'):
            excel_file.open('rb')
            try:
                if hasattr(excel_file, 'seek'):
                    excel_file.seek(0)
                workbook = load_workbook(excel_file)
                if hasattr(excel_file, 'seek'):
                    excel_file.seek(0)
                return workbook
            finally:
                if hasattr(excel_file, 'close'):
                    excel_file.close()

        if hasattr(excel_file, 'seek'):
            excel_file.seek(0)
        workbook = load_workbook(excel_file)
        if hasattr(excel_file, 'seek'):
            excel_file.seek(0)
        return workbook
    except Exception as exc:
        raise ValidationError('The uploaded file could not be parsed as Excel.') from exc


def _detect_headers(rows):
    for row_index, row in enumerate(rows):
        values = [value for value in row if value not in (None, '')]
        if values:
            last_used_index = max(index for index, value in enumerate(row) if value not in (None, ''))
            trimmed_row = row[: last_used_index + 1]
            return row_index, [value if value not in (None, '') else f'Column {index + 1}' for index, value in enumerate(trimmed_row)]
    raise ValidationError('No header row was found in the Excel file.')


def _serialize_cell(value):
    if value is None:
        return ''
    return value